from pathlib import Path
import re
from docx import Document
from docx.shared import Cm,Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH,WD_BREAK
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
src=Path("roman/Die_Glocken_von_Dornwacht.md").read_text(encoding="utf-8")
items=[]
for line in src.splitlines():
 s=line.strip()
 if not s or s.startswith("<!--"):continue
 if s.startswith(">"):s=re.sub(r"^>\s*(?:\*\*SKARRA\s*/\s*SPIELER:\*\*\s*)?","",s)
 if not s:continue
 if s.startswith("# "):continue
 if s.startswith("## "):items.append(("chapter",s[3:].strip()));continue
 if s.startswith("### "):items.append(("section",s[4:].strip()));continue
 if s in ("---","* * *","***"):items.append(("break",""));continue
 items.append(("text",s))
def clean(s):
 s=re.sub(r"\[([^\]]+)\]\([^)]+\)",r"\1",s)
 return re.sub(r"(?<!\\)[*_]+","",s).strip()
def dialog(s):return clean(s).startswith(("„","»","“",'"',"‚","– "))
out=[];pending=[]
def flush():
 if pending:out.append(("paragraph"," ".join(pending)));pending.clear()
for kind,s in items:
 if kind!="text" or dialog(s):
  flush();out.append((kind if kind!="text" else "paragraph",s));continue
 if pending and (sum(map(len,pending))+len(s)>410 or len(pending)>=4):flush()
 pending.append(s)
flush()
doc=Document();sec=doc.sections[0]
sec.page_width=Cm(12);sec.page_height=Cm(19)
sec.top_margin=Cm(1.7);sec.bottom_margin=Cm(1.75)
sec.left_margin=Cm(1.9);sec.right_margin=Cm(1.45);sec.gutter=Cm(.3)
sec.footer_distance=Cm(.85)
settings=doc.settings.element
mirror=OxmlElement("w:mirrorMargins");settings.append(mirror)
normal=doc.styles["Normal"];normal.font.name="Garamond";normal.font.size=Pt(10.5)
normal.paragraph_format.line_spacing=1.13
normal.paragraph_format.space_before=Pt(0);normal.paragraph_format.space_after=Pt(0)
normal.paragraph_format.first_line_indent=Cm(.42)
for name,size in [("BookChapter",15),("BookSection",12)]:
 st=doc.styles.add_style(name,WD_STYLE_TYPE.PARAGRAPH)
 st.base_style=normal;st.font.name="Garamond";st.font.size=Pt(size);st.font.bold=True
 st.paragraph_format.first_line_indent=Cm(0);st.paragraph_format.space_after=Pt(14)
 st.paragraph_format.keep_with_next=True
sec.different_first_page_header_footer=True
footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
r=footer.add_run();r.font.name="Garamond";r.font.size=Pt(9)
fld=OxmlElement("w:fldSimple");fld.set(qn("w:instr"),"PAGE");r._r.addnext(fld)
title=doc.add_paragraph();title.alignment=WD_ALIGN_PARAGRAPH.CENTER
title.paragraph_format.space_before=Pt(140)
r=title.add_run("DIE GLOCKEN\nVON DORNWACHT");r.bold=True;r.font.name="Garamond";r.font.size=Pt(22)
p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.space_before=Pt(24)
r=p.add_run("Ein Fantasyroman");r.italic=True;r.font.size=Pt(12)
p.add_run().add_break(WD_BREAK.PAGE)
chapter=0;start=True
for kind,s in out:
 if kind=="chapter":
  if s.startswith("Prolog"):heading="PROLOG\nDie Ratten"
  else:
   chapter+=1
   heading="KAPITEL "+str(chapter)+"\n"+re.sub(r"^Kapitel\s+\d+\s*[·:–-]?\s*","",s,flags=re.I)
  p=doc.add_paragraph(style="BookChapter");p.paragraph_format.page_break_before=True
  p.add_run(heading);start=True
 elif kind=="section":
  p=doc.add_paragraph(style="BookSection");p.add_run(clean(s));start=True
 elif kind=="break":
  if chapter==0:continue
  p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER
  p.paragraph_format.space_before=Pt(8);p.paragraph_format.space_after=Pt(8)
  p.paragraph_format.first_line_indent=Cm(0);p.add_run("✦");start=True
 else:
  t=clean(s)
  if not t:continue
  p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
  if start:p.paragraph_format.first_line_indent=Cm(0)
  p.add_run(t);start=False
update=OxmlElement("w:updateFields");update.set(qn("w:val"),"true");settings.append(update)
output=Path("export/Die_Glocken_von_Dornwacht.docx");output.parent.mkdir(exist_ok=True,parents=True)
doc.save(output)
print("Created",output,"chapters",chapter,"bytes",output.stat().st_size)
